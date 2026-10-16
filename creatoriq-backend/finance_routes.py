"""Milestone 3: revenue & sponsorship tracking, alerts/notifications, PDF/Excel reports."""
import io
import os
import random
from datetime import date, timedelta

from fastapi import APIRouter, Depends, Header, HTTPException, Response
from jose import jwt
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, Column, Date, DateTime, Float, ForeignKey, Integer, String, UniqueConstraint
from sqlalchemy.orm import Session
from sqlalchemy.sql import func

import models
from database import Base, get_db
from youtube_routes import Snapshot

router = APIRouter(prefix="/api", tags=["finance"])
SOURCES = ["Sponsorship", "Ad Revenue", "Affiliate Marketing", "Brand Collaboration", "Subscription"]
STATUSES = ["paid", "pending"]


class RevenueEntry(Base):
    __tablename__ = "revenue_entries"
    id = Column(Integer, primary_key=True)
    user_id = Column(Integer, ForeignKey("users.id"), index=True, nullable=False)
    source = Column(String, nullable=False)
    brand = Column(String, default="")
    platform = Column(String, default="YouTube")
    amount = Column(Float, nullable=False)
    date = Column(Date, nullable=False)
    status = Column(String, default="paid")
    notes = Column(String, default="")


class Notification(Base):
    __tablename__ = "notifications"
    id = Column(Integer, primary_key=True)
    user_id = Column(Integer, ForeignKey("users.id"), index=True, nullable=False)
    key = Column(String, nullable=False)
    kind = Column(String, default="info")
    message = Column(String, nullable=False)
    is_read = Column(Boolean, default=False)
    created_at = Column(DateTime(timezone=True), server_default=func.now())
    __table_args__ = (UniqueConstraint("user_id", "key"),)


class EntryIn(BaseModel):
    source: str
    brand: str = ""
    platform: str = "YouTube"
    amount: float = Field(gt=0)
    date: date
    status: str = "paid"
    notes: str = ""


def current_user(authorization: str = Header(None), db: Session = Depends(get_db)):
    try:
        email = jwt.decode(authorization.split()[1], os.environ["SECRET_KEY"], algorithms=["HS256"])["sub"]
    except Exception:
        raise HTTPException(401, "Not authenticated")
    user = db.query(models.User).filter_by(email=email).first()
    if not user or not user.is_active:
        raise HTTPException(401, "Not authenticated")
    return user


def need(*roles):
    """Route guard: only the listed roles may call the endpoint."""
    def dep(user=Depends(current_user)):
        if user.role not in roles:
            raise HTTPException(403, "Your role does not have access to this feature")
        return user
    return dep


def status_of(r):  # a pending payment past its date counts as overdue
    return "overdue" if r.status == "pending" and r.date < date.today() else r.status


def ser(r):
    return {"id": r.id, "source": r.source, "brand": r.brand, "platform": r.platform, "amount": r.amount,
            "date": r.date.isoformat(), "status": status_of(r), "notes": r.notes}


def check(body: EntryIn):
    if body.source not in SOURCES:
        raise HTTPException(422, f"source must be one of {SOURCES}")
    if body.status not in STATUSES:
        raise HTTPException(422, "status must be paid or pending")


def mine(db, user):
    return db.query(RevenueEntry).filter_by(user_id=user.id).order_by(RevenueEntry.date.desc(), RevenueEntry.id.desc()).all()


@router.get("/revenue")
def list_revenue(user=Depends(need("creator", "agency", "admin")), db: Session = Depends(get_db)):
    return [ser(r) for r in mine(db, user)]


@router.post("/revenue")
def add_revenue(body: EntryIn, user=Depends(need("creator", "agency", "admin")), db: Session = Depends(get_db)):
    check(body)
    r = RevenueEntry(user_id=user.id, **body.model_dump())
    db.add(r); db.commit(); db.refresh(r)
    return ser(r)


@router.put("/revenue/{rid}")
def edit_revenue(rid: int, body: EntryIn, user=Depends(need("creator", "agency", "admin")), db: Session = Depends(get_db)):
    check(body)
    r = db.query(RevenueEntry).filter_by(id=rid, user_id=user.id).first()
    if not r:
        raise HTTPException(404, "Entry not found")
    for k, v in body.model_dump().items():
        setattr(r, k, v)
    db.commit()
    return ser(r)


@router.delete("/revenue/{rid}")
def delete_revenue(rid: int, user=Depends(need("creator", "agency", "admin")), db: Session = Depends(get_db)):
    r = db.query(RevenueEntry).filter_by(id=rid, user_id=user.id).first()
    if not r:
        raise HTTPException(404, "Entry not found")
    db.delete(r); db.commit()
    return {"deleted": rid}


@router.delete("/revenue")
def clear_revenue(user=Depends(need("creator", "agency", "admin")), db: Session = Depends(get_db)):
    """Removes all of this user's income entries and their revenue/report alerts (e.g. after demo data)."""
    n = db.query(RevenueEntry).filter_by(user_id=user.id).delete()
    db.query(Notification).filter(Notification.user_id == user.id, Notification.kind.in_(["revenue", "report"])).delete(synchronize_session=False)
    db.commit()
    return {"deleted": n}


@router.post("/revenue/demo")
def demo_data(user=Depends(need("creator", "agency", "admin")), db: Session = Depends(get_db)):
    """Adds sample income so the charts can be demonstrated. Clearly a demo helper."""
    rnd, today = random.Random(42), date.today()
    brands = ["NordVPN", "Skillshare", "Squarespace", "Notion", "Audible", "Ridge Wallet"]
    for m in range(5, -1, -1):
        base = today.replace(day=15) - timedelta(days=30 * m)
        rows = [("Ad Revenue", "YouTube AdSense", rnd.randint(1800, 3200), "paid"),
                ("Subscription", "Channel memberships", rnd.randint(250, 600), "paid"),
                ("Affiliate Marketing", "Amazon Associates", rnd.randint(300, 900), "paid"),
                ("Sponsorship", rnd.choice(brands), rnd.randint(1500, 4500), "paid")]
        if m % 2 == 0:
            rows.append(("Brand Collaboration", rnd.choice(brands), rnd.randint(2000, 5000), "paid"))
        for s, b, a, st in rows:
            db.add(RevenueEntry(user_id=user.id, source=s, brand=b, amount=a, date=base, status=st))
    db.add(RevenueEntry(user_id=user.id, source="Sponsorship", brand="Squarespace", amount=3200, date=today + timedelta(days=9), status="pending"))
    db.add(RevenueEntry(user_id=user.id, source="Sponsorship", brand="Audible", amount=2400, date=today - timedelta(days=8), status="pending"))
    db.commit()
    return {"added": True}


def month_key(d):
    return d.strftime("%Y-%m")


def summarize(rows):
    today = date.today()
    paid = [r for r in rows if status_of(r) == "paid"]
    by_source = {s: 0.0 for s in SOURCES}
    months = {}
    for i in range(11, -1, -1):
        y, mo = today.year, today.month - i
        while mo < 1:
            mo += 12; y -= 1
        months[f"{y}-{mo:02d}"] = 0.0
    for r in paid:
        by_source[r.source] += r.amount
        if month_key(r.date) in months:
            months[month_key(r.date)] += r.amount
    keys = list(months)
    this_m, last_m = months[keys[-1]], months[keys[-2]]
    brands = {}
    for r in paid:
        if r.source in ("Sponsorship", "Brand Collaboration") and r.brand:
            brands[r.brand] = brands.get(r.brand, 0) + r.amount
    return {
        "total_earned": sum(r.amount for r in paid),
        "this_month": this_m, "last_month": last_m,
        "month_change_pct": round((this_m - last_m) / last_m * 100, 1) if last_m else None,
        "pending": sum(r.amount for r in rows if status_of(r) == "pending"),
        "overdue": sum(r.amount for r in rows if status_of(r) == "overdue"),
        "by_source": [{"source": k, "amount": round(v, 2)} for k, v in by_source.items() if v],
        "monthly": [{"month": k, "amount": round(v, 2)} for k, v in months.items()],
        "top_brand": max(brands, key=brands.get) if brands else None,
    }


@router.get("/revenue/summary")
def revenue_summary(user=Depends(need("creator", "agency", "admin")), db: Session = Depends(get_db)):
    return summarize(mine(db, user))


# ---------------- alerts ----------------
def refresh_alerts(db: Session, user):
    def add(key, kind, msg):
        if not db.query(Notification).filter_by(user_id=user.id, key=key).first():
            db.add(Notification(user_id=user.id, key=key, kind=kind, message=msg))

    rows = mine(db, user)
    for r in rows:
        if status_of(r) == "overdue":
            add(f"overdue-{r.id}", "revenue", f"Payment overdue: {r.brand or r.source} owes ${r.amount:,.0f} (due {r.date:%b %d})")
    total = sum(r.amount for r in rows if status_of(r) == "paid")
    crossed = [t for t in (1000, 5000, 10000, 25000, 50000, 100000) if total >= t]
    if crossed:
        add(f"rev-{crossed[-1]}", "revenue", f"Revenue milestone: you have earned over ${crossed[-1]:,} in total")
    week = sum(r.amount for r in rows if status_of(r) == "paid" and r.date >= date.today() - timedelta(days=7))
    iso = date.today().isocalendar()
    add(f"weekly-{iso[0]}-{iso[1]}", "report", f"Weekly summary: ${week:,.0f} earned in the last 7 days")

    for (cid,) in db.query(Snapshot.channel_id).distinct().all():
        snaps = db.query(Snapshot).filter_by(channel_id=cid).order_by(Snapshot.day.desc()).limit(9).all()
        if len(snaps) < 2:
            continue
        name = snaps[0].handle or cid
        if snaps[0].subs // 1_000_000 > snaps[1].subs // 1_000_000:
            add(f"subs-{cid}-{snaps[0].subs // 1_000_000}", "growth", f"{name} passed {snaps[0].subs // 1_000_000}M subscribers")
        deltas = [snaps[i].views - snaps[i + 1].views for i in range(len(snaps) - 1)]
        avg = sum(deltas[1:]) / len(deltas[1:]) if len(deltas) > 1 else 0
        if avg > 0 and deltas[0] < avg * 0.5:
            add(f"drop-{cid}-{snaps[0].day}", "warning", f"{name}: daily views fell below half of the recent average")
        if avg > 0 and deltas[0] > avg * 2:
            add(f"surge-{cid}-{snaps[0].day}", "growth", f"{name}: daily views are more than double the recent average")
    db.commit()


@router.get("/notifications")
def notifications(user=Depends(current_user), db: Session = Depends(get_db)):
    refresh_alerts(db, user)
    items = db.query(Notification).filter_by(user_id=user.id).order_by(Notification.created_at.desc(), Notification.id.desc()).limit(50).all()
    return {"unread": sum(1 for n in items if not n.is_read),
            "items": [{"id": n.id, "kind": n.kind, "message": n.message, "is_read": n.is_read, "created_at": n.created_at.isoformat() if n.created_at else None} for n in items]}


@router.post("/notifications/read-all")
def read_all(user=Depends(current_user), db: Session = Depends(get_db)):
    db.query(Notification).filter_by(user_id=user.id, is_read=False).update({"is_read": True})
    db.commit()
    return {"ok": True}


# ---------------- exports ----------------
def build_file(fmt, title, lines, cols, rows):
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "Report"
        ws.append([title]); ws["A1"].font = Font(bold=True, size=14)
        for line in lines:
            ws.append([line])
        ws.append([]); ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1A6DFF")
        for r in rows:
            ws.append(r)
        for i in range(len(cols)):
            ws.column_dimensions[chr(65 + i)].width = 20
        buf = io.BytesIO(); wb.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO(); st = getSampleStyleSheet()
    els = [Paragraph(title, st["Title"])] + [Paragraph(x, st["Normal"]) for x in lines] + [Spacer(1, 12)]
    t = Table([cols] + [[str(x) for x in r] for r in rows], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a6dff")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f4fb")])]))
    els.append(t)
    SimpleDocTemplate(buf, pagesize=A4).build(els)
    return buf.getvalue(), "application/pdf"


@router.get("/reports/export")
def export_report(kind: str, fmt: str = "xlsx", handle: str = "", user=Depends(current_user), db: Session = Depends(get_db)):
    if fmt not in ("xlsx", "pdf") or kind not in ("revenue", "growth"):
        raise HTTPException(422, "kind must be revenue or growth; fmt must be xlsx or pdf")
    if kind == "revenue" and user.role == "marketing":
        raise HTTPException(403, "Your role cannot access revenue reports")
    if kind == "revenue":
        rows, s = mine(db, user), None
        s = summarize(rows)
        lines = [f"Generated {date.today():%d %b %Y}", f"Total earned: ${s['total_earned']:,.2f}", f"This month: ${s['this_month']:,.2f}",
                 f"Pending: ${s['pending']:,.2f}   Overdue: ${s['overdue']:,.2f}"]
        data = [[r.date.isoformat(), r.source, r.brand, r.platform, round(r.amount, 2), status_of(r)] for r in rows]
        content, mime = build_file(fmt, "CreatorIQ Revenue Report", lines, ["Date", "Source", "Brand", "Platform", "Amount ($)", "Status"], data)
    else:
        h = handle.strip().lower()
        h = h if h.startswith("@") else "@" + h
        snaps = db.query(Snapshot).filter(func.lower(Snapshot.handle) == h).order_by(Snapshot.day).all()
        if not snaps:
            raise HTTPException(404, "No tracked history for that channel yet. Search it on the YouTube page first.")
        data = []
        for i, sn in enumerate(snaps):
            data.append([sn.day.isoformat(), sn.subs, sn.views, sn.videos, sn.views - snaps[i - 1].views if i else "", sn.subs - snaps[i - 1].subs if i else ""])
        lines = [f"Channel {h}", f"Tracked days: {len(snaps)}", f"Latest subscribers: {snaps[-1].subs:,}   Latest views: {snaps[-1].views:,}"]
        content, mime = build_file(fmt, "CreatorIQ Growth Report", lines, ["Date", "Subscribers", "Views", "Videos", "Views change", "Subs change"], data)
    return Response(content, media_type=mime, headers={"Content-Disposition": f'attachment; filename="creatoriq-{kind}-report.{fmt}"'})